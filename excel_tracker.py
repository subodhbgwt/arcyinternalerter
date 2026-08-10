"""Append newly-found roles to a running Excel tracker (job_tracker.xlsx).

One row per role, deduped by the job's stable id (stored in a hidden first
column). Existing rows are never rewritten, so any Status / Notes you type in
Excel survive across runs. Called from tracker.py after each digest.
"""
from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

log = logging.getLogger("arcy-alerter.xlsx")

HEADERS = ["ID", "Date found", "Company", "Role", "Location",
           "Type", "Track", "Status", "Link", "Materials", "Notes"]

STATUS_CHOICES = ["New", "Interested", "Applied", "Interview",
                  "Offer", "Rejected", "Skip"]

# Column widths (by header index) for readability.
_WIDTHS = [0, 12, 20, 46, 26, 15, 12, 12, 10, 34, 26]


def _try_import():
    try:
        import openpyxl  # noqa: F401
        return True
    except ImportError:
        log.warning("openpyxl not installed — Excel tracker skipped. "
                    "Run: pip install openpyxl")
        return False


def _classify(job: dict) -> str:
    try:
        from filters import classify_track
        return classify_track(job.get("title", ""), job.get("department", "") or "")
    except Exception:
        return ""


def append_jobs(jobs: list[dict], path: str | Path) -> int:
    """Append any jobs whose id isn't already in the sheet. Returns count added.
    Never overwrites existing rows (preserves your Status/Notes edits)."""
    if not jobs or not _try_import():
        return 0

    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    path = Path(path)

    if path.exists():
        try:
            wb = load_workbook(path)
            ws = wb.active
        except Exception as exc:
            log.warning("Couldn't open %s (%s) — is it open in Excel? "
                        "Skipping tracker write this run.", path.name, exc)
            return 0
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "Jobs"
        ws.append(HEADERS)
        hdr_fill = PatternFill("solid", fgColor="1F4E78")
        for i, _ in enumerate(HEADERS, start=1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
            c.alignment = Alignment(vertical="center")
            if _WIDTHS[i - 1]:
                ws.column_dimensions[c.column_letter].width = _WIDTHS[i - 1]
        ws.column_dimensions["A"].hidden = True   # hide the id column
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(HEADERS)).column_letter}1"

    # Existing ids (column A), so we never double-add.
    seen = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        if row and row[0]:
            seen.add(str(row[0]))

    today = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    added = 0
    for j in jobs:
        jid = str(j.get("id", ""))
        if not jid or jid in seen:
            continue
        seen.add(jid)
        ws.append([
            jid, today,
            j.get("company", ""), j.get("title", ""), j.get("location", ""),
            j.get("_label", ""), _classify(j), "New",
            "apply", "", "",
        ])
        link_cell = ws.cell(row=ws.max_row, column=9)
        url = j.get("url", "")
        if url:
            link_cell.hyperlink = url
            link_cell.font = Font(color="0563C1", underline="single")
        added += 1

    if not added:
        return 0

    # Status dropdown across the data range.
    dv = DataValidation(type="list",
                        formula1='"' + ",".join(STATUS_CHOICES) + '"',
                        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"H2:H{ws.max_row}")

    try:
        wb.save(path)
    except PermissionError:
        log.warning("Couldn't save %s — it's open in Excel. "
                    "Close it; the next run will catch up.", path.name)
        return 0
    except Exception as exc:
        log.warning("Tracker save failed: %s", exc)
        return 0

    log.info("Excel tracker: +%d row(s) -> %s", added, path.name)
    return added
