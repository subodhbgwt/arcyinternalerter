"""Application factory. Prepares packets, never submits.

For every row in job_tracker.xlsx marked Status = "Interested" that has no
Materials yet, it prepares an application packet folder:

    applications/<Company>__<Role>/
        apply.txt          # role, link, and a submit checklist
        CV_<track>.pdf      # the right CV variant (finance/quant/tech), copied in
        cover_letter.md     # a per-role brief + PENDING marker until the letter is written

Then it writes the folder path into the tracker's Materials column and prints
the packets that still need a cover letter.

It deliberately does NOT submit anything. You review each packet and apply
yourself. Quality control matters at these firms (many have reapplication
cooldowns), and account/login/CAPTCHA steps aren't safe to automate.

Usage:
    python apply_factory.py            # prep packets for "Interested" rows
    python apply_factory.py --list     # just list packets needing a letter
"""
from __future__ import annotations

import argparse
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).parent
TRACKER = ROOT / "job_tracker.xlsx"
APPS_DIR = ROOT / "applications"

# Where Subodh's CV PDFs live, and which variant fits each track.
CV_DIR = Path(r"C:\Users\subodh\Downloads\job")
CV_BY_TRACK = {
    "Quant":       "qcurriculumvitae.pdf",
    "Finance/IB":  "fcurriculumvitae.pdf",
    "Consulting":  "fcurriculumvitae.pdf",
    "Tech/SWE":    "tcurriculumvitae.pdf",
    "Programme":   "fcurriculumvitae.pdf",
}
CV_DEFAULT = "fcurriculumvitae.pdf"

PENDING = "<!-- LETTER PENDING -->"


def _slug(*parts: str) -> str:
    s = "__".join(parts)
    s = re.sub(r"[^\w\s-]", "", s).strip()
    s = re.sub(r"[\s]+", "_", s)
    return s[:80]


def _headers(ws) -> dict:
    return {ws.cell(1, c).value: c
            for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}


def _ensure_materials_column(ws) -> None:
    """Insert a Materials column before Notes if an older sheet lacks it."""
    hdr = _headers(ws)
    if "Materials" in hdr:
        return
    notes_c = hdr.get("Notes", ws.max_column + 1)
    ws.insert_cols(notes_c)
    ws.cell(1, notes_c).value = "Materials"
    try:
        from openpyxl.styles import Font, PatternFill
        c = ws.cell(1, notes_c)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        ws.column_dimensions[c.column_letter].width = 34
    except Exception:
        pass


def _write_packet(company, role, location, link, track, folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)

    # 1) copy the right CV variant
    cv_name = CV_BY_TRACK.get(track, CV_DEFAULT)
    src = CV_DIR / cv_name
    cv_note = ""
    if src.exists():
        shutil.copy2(src, folder / f"CV_{track.replace('/', '-')}.pdf")
    else:
        cv_note = f"(!) CV file not found: {src}, attach it manually.\n"

    # 2) apply.txt: link and checklist
    (folder / "apply.txt").write_text(
        f"COMPANY : {company}\n"
        f"ROLE    : {role}\n"
        f"LOCATION: {location}\n"
        f"TRACK   : {track}\n"
        f"APPLY   : {link}\n\n"
        f"{cv_note}"
        "CHECKLIST before you submit:\n"
        "  [ ] Read cover_letter.md, tweak the one line that needs your voice\n"
        "  [ ] Tailor the CV top line if the role is oddly specific\n"
        "  [ ] Confirm work-authorisation answer (EU citizen / no sponsorship)\n"
        "  [ ] Check the deadline / that it's still open\n"
        "  [ ] Answer any 'why us' box (reuse the cover letter's middle para)\n"
        "  [ ] Submit, then set Status = Applied in the tracker\n",
        encoding="utf-8",
    )

    # 3) cover_letter.md: brief plus the pending marker
    (folder / "cover_letter.md").write_text(
        f"# Cover letter: {company} · {role}\n\n"
        f"- Location: {location}\n"
        f"- Track: {track}\n"
        f"- Link: {link}\n\n"
        f"{PENDING}\n\n"
        "_Angle to use: engineer/builder who has already built the analytical "
        "tools this field runs on (DCF, options pricer, trading bot, AI research "
        "agent). Humanised, no AI tells, British spelling, true claims only._\n",
        encoding="utf-8",
    )


def run(list_only: bool = False) -> int:
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl not installed. Run: pip install openpyxl")

    if not TRACKER.exists():
        sys.exit(f"No tracker at {TRACKER}. Run a digest first.")

    try:
        wb = load_workbook(TRACKER)
    except Exception as exc:
        sys.exit(f"Couldn't open {TRACKER.name} ({exc}). Close it in Excel first.")
    ws = wb.active

    if not list_only:
        _ensure_materials_column(ws)
    hdr = _headers(ws)
    ci = {name: hdr[name] for name in
          ("Company", "Role", "Location", "Status", "Link", "Track", "Materials")
          if name in hdr}

    prepared, pending = [], []
    for r in range(2, ws.max_row + 1):
        status = (ws.cell(r, ci["Status"]).value or "").strip().lower()
        if status != "interested":
            continue
        company = ws.cell(r, ci["Company"]).value or "Company"
        role = ws.cell(r, ci["Role"]).value or "Role"
        location = ws.cell(r, ci["Location"]).value or ""
        track = ws.cell(r, ci["Track"]).value or "Programme"
        link_cell = ws.cell(r, ci["Link"])
        link = link_cell.hyperlink.target if link_cell.hyperlink else (link_cell.value or "")
        materials = ws.cell(r, ci["Materials"]).value if "Materials" in ci else None

        folder = APPS_DIR / _slug(str(company), str(role))
        letter = folder / "cover_letter.md"
        needs_letter = (not letter.exists()) or ("LETTER PENDING" in letter.read_text(encoding="utf-8"))

        if not list_only and not materials:
            _write_packet(company, role, location, link, track, folder)
            ws.cell(r, ci["Materials"]).value = str(folder.relative_to(ROOT))
            prepared.append(folder.name)
        if needs_letter:
            pending.append({"folder": str(folder), "company": company,
                            "role": role, "location": location, "track": track,
                            "link": link})

    if not list_only:
        try:
            wb.save(TRACKER)
        except PermissionError:
            sys.exit(f"Couldn't save {TRACKER.name}. Close it in Excel and run again.")
        print(f"Prepared {len(prepared)} packet(s): {', '.join(prepared) or '(none new)'}")

    print(f"\n{len(pending)} packet(s) still need a cover letter:")
    for p in pending:
        print(f"  - {p['company']} · {p['role']}  [{p['track']}]  -> {p['folder']}")
    if pending:
        print("\nWrite the letter in each cover_letter.md, then apply.")
    return 0


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Prepare review-gated application packets.")
    ap.add_argument("--list", action="store_true",
                    help="Only list packets needing a cover letter; don't create/prep.")
    args = ap.parse_args()
    sys.exit(run(list_only=args.list))
